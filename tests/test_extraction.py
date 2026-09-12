import app


def test_extract_records_smart_954_counts(agent954_files):
    old_file, new_file = agent954_files
    old_data, old_dupes = app.extract_records_smart(old_file, card_type="new")
    new_data, new_dupes = app.extract_records_smart(new_file, card_type="new")
    assert len(old_data) == 254
    assert len(new_data) == 402
    assert old_dupes == []
    assert new_dupes == []


def test_extract_records_smart_921_counts(agent921_files):
    old_file, new_file = agent921_files
    old_data, _ = app.extract_records_smart(old_file, card_type="new")
    new_data, _ = app.extract_records_smart(new_file, card_type="new")
    assert len(old_data) == 406
    assert len(new_data) == 422


def test_extract_records_smart_handles_multi_sheet_without_repeated_header(agent954_files):
    """ملف 954 الإكسل مقسوم على شيتين (Table 1 له صف عناوين، Table 2
    بيانات فقط بدون تكرار العناوين) — هذا بالضبط ما يثبت إن last_role_map
    يشتغل صح عبر أوراق متعددة."""
    old_file, _ = agent954_files
    old_data, _ = app.extract_records_smart(old_file, card_type="new")
    assert len(old_data) > 200  # لو فشل بالورقة الثانية كان العدد يطلع أقل بكثير


def test_preview_columns_detects_expected_roles_xlsx(agent954_files):
    old_file, _ = agent954_files
    preview = app.preview_columns_for_file(old_file)
    assert preview is not None
    assert "اسم رب الأسرة" in preview["detected"]
    assert "الأفراد المستحقة" in preview["detected"]
    assert "الأفراد المحجوبين" in preview["detected"]
    assert len(preview["sample_records"]) == 2


def test_preview_columns_detects_expected_roles_docx(agent954_files):
    _, new_file = agent954_files
    preview = app.preview_columns_for_file(new_file)
    assert preview is not None
    assert "اسم رب الأسرة" in preview["detected"]
    assert "الأفراد الكلية" in preview["detected"]


def test_preview_columns_handles_merged_single_cell_header():
    """بعض الملفات القديمة تكتب كل أسماء الأعمدة بخلية واحدة (بدل عمود
    منفصل لكل عنوان)، مع عمود فارغ مخفي بين الاسم والبطاقات وترتيب أعمدة
    فعلي (محجوب, مستحق, كلي, اسم, ..., بطاقات, تسلسل) لا يطابق ترتيب النص
    بخلية العنوان. preview_columns_for_file لازم يتعرف على هذا الشكل
    (merged_header=True) بدل ما يرجع None ويطلع تحذير "ما فيه جدول واضح"
    المضلل، مع عيّنة صحيحة القيم مبنية عبر المحرك الاحتياطي الموثوق."""
    import io
    import openpyxl

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.append(["نوع الوكالة          FOOD              رقم المركز                       625"])
    ws.append(["رقم الوكالة         038526            اسم الوكيل                     علي حيدر حسين"])
    ws.append(["ت                    رقم البطاقة                  رقم البطاقة القديم                                            اسم رب الاسرة                                       الافراد الكلية       الافراد المستحقة   الافراد المحجوبين"])
    ws.append([])
    # [محجوب, مستحق, كلي, اسم, فارغ, بطاقة_جديدة, بطاقة_قديمة, تسلسل]
    ws.append([0, 8, 8, "وهاب جرد كاظم", None, 46774, 4454273, 1])
    ws.append([0, 2, 2, "صادق رزاق جبار ترابي", None, 47899, 4351147, 2])
    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    buf.name = "merged_header.xlsx"

    preview = app.preview_columns_for_file(buf)
    assert preview is not None
    assert preview.get("merged_header") is True
    assert "اسم رب الأسرة" in preview["detected"]
    assert len(preview["sample_records"]) == 2
    assert preview["sample_records"][0]["اسم رب الأسرة"] == "وهاب جرد كاظم"
    assert preview["sample_records"][0]["الأفراد الكلية"] == 8
    assert preview["sample_records"][0]["الأفراد المستحقة"] == 8
    assert preview["sample_records"][0]["الأفراد المحجوبين"] == 0

    buf.seek(0)
    data, dupes = app.extract_clean_records(buf, card_type="old")
    assert len(data) == 2
    assert dupes == []
    rec = data["46774"]
    assert rec["name"] == "وهاب جرد كاظم"
    assert rec["total"] == 8 and rec["eligible"] == 8 and rec["withheld"] == 0
    assert rec["alt_card"] == "4454273"


def test_preview_columns_returns_none_for_file_with_no_clear_table():
    """ملف نصي عشوائي بدون أي جدول/عناوين معروفة — لازم يرجع None بدل ما
    يرمي استثناء أو يخمّن أعمدة عشوائية."""
    import io
    buf = io.BytesIO()
    from docx import Document
    doc = Document()
    doc.add_paragraph("هذا ملف بدون أي جدول بيانات إطلاقاً.")
    doc.save(buf)
    buf.seek(0)
    buf.name = "empty.docx"
    assert app.preview_columns_for_file(buf) is None
