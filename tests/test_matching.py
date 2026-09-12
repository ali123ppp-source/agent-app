import app


def test_dual_key_matching_921(agent921_files):
    old_file, new_file = agent921_files
    old_data, new_data, card_col_name, old_dupes, new_dupes = app.extract_matched_by_either_card(
        app.extract_records_smart, old_file, new_file
    )
    common = set(old_data.keys()) & set(new_data.keys())
    assert len(old_data) == 406
    assert len(new_data) == 422
    assert len(common) == 396
    assert card_col_name == "رقم البطاقة"
    assert old_dupes == []
    assert new_dupes == []


def test_dual_key_matching_954(agent954_files):
    old_file, new_file = agent954_files
    old_data, new_data, card_col_name, _, _ = app.extract_matched_by_either_card(
        app.extract_records_smart, old_file, new_file
    )
    common = set(old_data.keys()) & set(new_data.keys())
    # كل عوائل الإكسل القديم (254) لهم تطابق بالوورد الحديث — صفر محذوفات
    assert len(common) == 254


def test_dual_key_matching_with_mismatched_files_finds_no_matches(agent921_mismatched_files):
    """هذا زوج ملفات مش متناظر فعلياً (لقطتان مختلفتان) — التطابق صفر هو
    السلوك الصحيح هنا، مو خطأ. الاختبار يوثّق هذا السلوك المتوقع صراحة
    بدل ما يُكتشف بالصدفة بالإنتاج زي ما صار فعلياً."""
    old_file, new_file = agent921_mismatched_files
    old_data, new_data, _, _, _ = app.extract_matched_by_either_card(
        app.extract_records_smart, old_file, new_file
    )
    common = set(old_data.keys()) & set(new_data.keys())
    assert len(common) == 0


def test_merge_records_by_either_card_matches_via_alt_when_primary_differs():
    """اختبار وحدة مركّز لدالة الدمج نفسها بمعزل عن قراءة الملفات: عائلة
    نفس البطاقة الحديثة بس بطاقة قديمة مختلفة الشكل بالملفين لازم توصف
    كمتطابقة عبر alt_card، مو تُحسب مضافة ومحذوفة بالخطأ."""
    old_data = {
        "1111": {"name": "a", "total": 3, "eligible": 3, "withheld": 0, "alt_card": "9999999"},
    }
    new_data = {
        "2222": {"name": "a", "total": 3, "eligible": 2, "withheld": 1, "alt_card": "9999999"},
    }
    unified_old, unified_new = app.merge_records_by_either_card(old_data, new_data)
    assert set(unified_old.keys()) == {"1111"}
    assert unified_new["1111"]["eligible"] == 2  # جاب سجل new الصحيح عبر alt_card


def test_merge_records_by_either_card_no_double_counting_on_ambiguous_alt():
    """لو سجلين قدامى عندهم نفس alt_card (تصادم نادر)، ما يصير أي سجل
    جديد يتحسب مرتين كمتطابق مع الاثنين."""
    old_data = {
        "1111": {"name": "a", "total": 1, "eligible": 1, "withheld": 0, "alt_card": "5555555"},
        "1112": {"name": "b", "total": 2, "eligible": 2, "withheld": 0, "alt_card": "5555555"},
    }
    new_data = {
        "2222": {"name": "a", "total": 1, "eligible": 1, "withheld": 0, "alt_card": "5555555"},
    }
    unified_old, unified_new = app.merge_records_by_either_card(old_data, new_data)
    matched_keys = [k for k in unified_old if k in unified_new]
    assert len(matched_keys) == 1  # مو الاثنين


def test_extract_records_smart_normalizes_arabic_indic_digits_for_matching():
    """بطاقة مكتوبة بأرقام هندية-عربية بملف وبأرقام لاتينية بالثاني لازم
    تنعتبر نفس البطاقة (نفس المفتاح) بعد التطبيع، مو مضافة ومحذوفة بالغلط."""
    import io
    from docx import Document

    def make_docx(card_digits):
        buf = io.BytesIO()
        doc = Document()
        table = doc.add_table(rows=2, cols=4)
        header = table.rows[0].cells
        header[0].text = "ت"
        header[1].text = "رقم البطاقة"
        header[2].text = "اسم رب الأسرة"
        header[3].text = "الأفراد المستحقة"
        row = table.rows[1].cells
        row[0].text = "1"
        row[1].text = card_digits
        row[2].text = "احمد علي"
        row[3].text = "3"
        doc.save(buf)
        buf.seek(0)
        buf.name = "f.docx"
        return buf

    latin_file = make_docx("123456")
    indic_file = make_docx("١٢٣٤٥٦")

    latin_data, _ = app.extract_records_smart(latin_file, card_type="old")
    indic_data, _ = app.extract_records_smart(indic_file, card_type="old")
    assert set(latin_data.keys()) == set(indic_data.keys()) == {"123456"}


def test_smart_fallback_is_per_file_not_shared(agent921_files):
    """لو ملف واحد بس يفشل بالمحرك الذكي (عناوين مدمجة بخلية وحدة مثلاً)،
    ما لازم هذا يسحب الملف الثاني (اللي ناجح تماماً بالذكي) للمحرك القديم
    معاه — سيناريو فعلي شوهد يسبب تسرب رقم بطاقة لعمود عدد بالملف الثاني
    السليم أصلاً. agent921_new.docx يشتغل صح بالذكي وبيه فساد معروف لو
    قرأناه بالمحرك القديم (agent921_mismatched fixtures)، فنتأكد هنا إنه
    ما ينزل للقديم أبداً طالما الذكي نجح فيه، بغض النظر عن حال الملف الثاني."""
    old_file, new_file = agent921_files
    import app as app_module
    original_smart = app_module.extract_records_smart
    original_clean = app_module.extract_clean_records
    calls = {"smart_new": 0, "clean_new": 0}

    def patched_smart(file_obj, card_type="old"):
        if getattr(file_obj, "name", "") == old_file.name:
            return {}, []  # يحاكي فشل الذكي بملف old فقط
        calls["smart_new"] += 1
        return original_smart(file_obj, card_type=card_type)

    def patched_clean(file_obj, card_type="old"):
        if getattr(file_obj, "name", "") == new_file.name:
            calls["clean_new"] += 1
        return original_clean(file_obj, card_type=card_type)

    app_module.extract_records_smart = patched_smart
    app_module.extract_clean_records = patched_clean
    try:
        new_data, new_dupes, new_used_fallback = app_module._extract_with_smart_fallback(new_file, "old")
        old_data, old_dupes, old_used_fallback = app_module._extract_with_smart_fallback(old_file, "old")
    finally:
        app_module.extract_records_smart = original_smart
        app_module.extract_clean_records = original_clean

    assert new_used_fallback is False, "الملف الثاني ما لازم يهبط للمحرك القديم لأن الذكي نجح فيه"
    assert calls["clean_new"] == 0, "extract_clean_records ما لازم يُستدعى أصلاً على الملف السليم"
    assert calls["smart_new"] == 1
    assert len(new_data) == 422  # نفس نتيجة المحرك الذكي المعروفة لهذا الملف
    assert old_used_fallback is True  # الملف old فعلاً فشل بالذكي (محاكاة)


def test_extract_records_smart_reports_duplicate_card_within_same_file():
    """بطاقتين بنفس الرقم بنفس الملف: أول ظهور يُعتمد، الثاني يُسجَّل
    كمكرر بدل ما يُكتب فوق الأول بصمت (فقدان عائلة كاملة من الحساب)."""
    import io
    from docx import Document

    buf = io.BytesIO()
    doc = Document()
    table = doc.add_table(rows=3, cols=4)
    header = table.rows[0].cells
    header[0].text = "ت"
    header[1].text = "رقم البطاقة"
    header[2].text = "اسم رب الأسرة"
    header[3].text = "الأفراد المستحقة"
    r1 = table.rows[1].cells
    r1[0].text, r1[1].text, r1[2].text, r1[3].text = "1", "111111", "احمد علي", "3"
    r2 = table.rows[2].cells
    r2[0].text, r2[1].text, r2[2].text, r2[3].text = "2", "111111", "محمد كريم", "4"
    doc.save(buf)
    buf.seek(0)
    buf.name = "f.docx"

    data, duplicates = app.extract_records_smart(buf, card_type="old")
    assert len(data) == 1
    assert data["111111"]["name"] == "احمد علي"  # أول ظهور محفوظ
    assert duplicates == [("111111", "محمد كريم")]
