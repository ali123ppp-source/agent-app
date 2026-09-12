"""اختبارات ميزة مقارنة عدة أزواج ملفات دفعة وحدة: تحديد أي ملف يرتبط
بأي ملف آخر تلقائياً حسب تطابق البيانات الفعلية (رقم البطاقة)، وليس
أسماء الملفات أو ترتيب الرفع أو حتى امتدادها (xlsx/docx)."""
import io

import app


def test_auto_pair_files_by_content_matches_correct_agents_despite_shuffled_order(agent954_files, agent921_files):
    """رفع ملفات وكيلين معاً بترتيب مقصود مبعثر — النظام لازم يربط كل
    وكيل بملفه الصحيح حسب تطابق أرقام البطاقات الفعلي، مو حسب ترتيب
    الرفع ولا نوع الملف."""
    old_954, new_954 = agent954_files
    old_921, new_921 = agent921_files

    files = [old_921, old_954, new_954, new_921]

    pairs, unmatched = app.auto_pair_files_by_content(files)

    assert unmatched == []
    assert len(pairs) == 2

    pair_map = {}
    for a, b, overlap in pairs:
        pair_map[frozenset((a.name, b.name))] = overlap

    assert frozenset((old_954.name, new_954.name)) in pair_map
    assert frozenset((old_921.name, new_921.name)) in pair_map
    # كل زوج صحيح لازم يكون له تطابق فعلي بعشرات البطاقات، مو صدفة رقم أو رقمين
    assert pair_map[frozenset((old_954.name, new_954.name))] > 50
    assert pair_map[frozenset((old_921.name, new_921.name))] > 50


def _make_simple_docx(rows, name=None):
    from docx import Document
    doc = Document()
    table = doc.add_table(rows=1 + len(rows), cols=4)
    header = table.rows[0].cells
    header[0].text, header[1].text, header[2].text, header[3].text = "ت", "رقم البطاقة", "اسم رب الأسرة", "الأفراد المستحقة"
    for i, (card, name_) in enumerate(rows, start=1):
        cells = table.rows[i].cells
        cells[0].text, cells[1].text, cells[2].text, cells[3].text = str(i), card, name_, "2"
    buf = io.BytesIO()
    doc.save(buf)
    buf.seek(0)
    buf.name = name or f"f{id(rows)}.docx"
    return buf


def test_auto_pair_files_by_content_leaves_zero_overlap_file_unmatched():
    """ملف ما عنده ولا بطاقة مشتركة مع أي ملف بالمجموعة لازم يرجع ضمن
    unmatched، مو يترابط غلط بأقرب شي متوفر."""
    matching_a = _make_simple_docx([("1001", "احمد علي"), ("1002", "محمد كريم")], "a.docx")
    matching_b = _make_simple_docx([("1001", "احمد علي"), ("1002", "محمد كريم")], "b.docx")
    lonely = _make_simple_docx([("9999999", "شخص غريب تماما")], "lonely.docx")

    pairs, unmatched = app.auto_pair_files_by_content([matching_a, lonely, matching_b])

    assert len(pairs) == 1
    assert {pairs[0][0].name, pairs[0][1].name} == {"a.docx", "b.docx"}
    assert pairs[0][2] == 2
    assert unmatched == [lonely]


def _make_simple_xlsx(rows, name):
    import openpyxl
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.append(["ت", "رقم البطاقة", "اسم رب الأسرة", "الأفراد المستحقة"])
    for i, (card, name_) in enumerate(rows, start=1):
        ws.append([i, card, name_, 2])
    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    buf.name = name
    return buf


def test_auto_pair_files_by_content_works_with_same_extension_only():
    """يشتغل حتى لو كل الملفات المرفوعة من نفس النوع (كلها xlsx مثلاً) —
    المطابقة تعتمد على البيانات الفعلية مو الامتداد."""
    a1 = _make_simple_xlsx([("1111", "احمد"), ("2222", "محمد")], "old1.xlsx")
    a2 = _make_simple_xlsx([("1111", "احمد"), ("2222", "محمد")], "new1.xlsx")
    b1 = _make_simple_xlsx([("3333", "علي"), ("4444", "حسين")], "old2.xlsx")
    b2 = _make_simple_xlsx([("3333", "علي"), ("4444", "حسين")], "new2.xlsx")

    pairs, unmatched = app.auto_pair_files_by_content([a1, b1, b2, a2])

    assert unmatched == []
    assert len(pairs) == 2
    pair_names = {frozenset((p[0].name, p[1].name)) for p in pairs}
    assert frozenset(("old1.xlsx", "new1.xlsx")) in pair_names
    assert frozenset(("old2.xlsx", "new2.xlsx")) in pair_names


def test_file_key_set_includes_alt_card():
    records = {
        "123": {"name": "a", "alt_card": "9999"},
        "456": {"name": "b", "alt_card": ""},
    }
    keys = app._file_key_set(records)
    assert keys == {"123", "9999", "456"}
